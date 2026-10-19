"""Render a tailored resume .docx from structured content.

Layout matches Matthew's established truth-first resume style
(Calibri, navy/blue accents, one page, ATS-friendly).
"""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(31, 54, 80)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(84, 84, 84)
BLACK = RGBColor(0, 0, 0)

# Personal lines live in config.json (untracked); the fallbacks are placeholders.
try:
    _cfg = json.loads((Path(__file__).parent / "config.json").read_text(encoding="utf-8"))
except (OSError, ValueError):
    _cfg = {}
CONTACT = _cfg.get("resume_contact",
                   "you@example.com | +1 (555) 555-0100 | linkedin.com/in/your-handle")
CITY = _cfg.get("resume_city", "Your City, ON")
EDUCATION = _cfg.get("resume_education",
                     "Your Degree - Your University | Graduated 20XX")


def _font(run, size, color=BLACK, bold=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def _para(p, before=0, after=2, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align


def _rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    borders.append(bottom)


def _heading(doc, text):
    p = doc.add_paragraph()
    _para(p, before=5, after=1)
    r = p.add_run(text.upper())
    _font(r, 9.1, BLUE, bold=True)
    _rule(p)


def _text(doc, text, size=9.05, bold=False, color=BLACK, after=2):
    p = doc.add_paragraph()
    _para(p, after=after)
    r = p.add_run(text)
    _font(r, size, color, bold=bold)


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    _para(p, after=1)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(text)
    _font(r, 8.8)


def build_cover_letter(body: str, out_docx: Path, date_str: str | None = None) -> Path:
    """Render a cover letter as a document in the same style as the resume.
    `body` is the plain-text letter (greeting through sign-off, paragraphs
    separated by blank lines) — used verbatim, nothing added or invented."""
    from datetime import datetime

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    _para(p, after=0)
    r = p.add_run("MATTHEW MCCUSKER")
    _font(r, 18, NAVY, bold=True)

    p = doc.add_paragraph()
    _para(p, after=2)
    r = p.add_run(CITY + "  ·  " + CONTACT.replace(" | ", "  ·  "))
    _font(r, 9.5, MUTED)
    _rule(p)

    p = doc.add_paragraph()
    _para(p, before=14, after=12)
    r = p.add_run(date_str or datetime.now().strftime("%B %d, %Y").replace(" 0", " "))
    _font(r, 11)

    paragraphs = [blk.strip() for blk in body.replace("\r", "").split("\n\n") if blk.strip()]
    for i, blk in enumerate(paragraphs):
        p = doc.add_paragraph()
        _para(p, after=4 if i == len(paragraphs) - 2 else 10)
        # keep intentional single line breaks (e.g. "Sincerely,\nMatthew McCusker")
        lines = blk.split("\n")
        for j, line in enumerate(lines):
            r = p.add_run(line.strip())
            _font(r, 11)
            if j < len(lines) - 1:
                r.add_break()

    doc.core_properties.author = "Matthew McCusker"
    doc.save(out_docx)
    return out_docx


def build_resume(data: dict, out_docx: Path) -> Path:
    """data keys: title_line, location_line, profile, role_match [{label,text}],
    experience [{title,company,dates,bullets[]}], education_note, keywords."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(9.05)

    p = doc.add_paragraph()
    _para(p, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("MATTHEW MCCUSKER")
    _font(r, 18.5, NAVY, bold=True)

    p = doc.add_paragraph()
    _para(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(data["title_line"])
    _font(r, 10.55, BLUE, bold=True)

    p = doc.add_paragraph()
    _para(p, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(data["location_line"])
    _font(r, 8.65)

    p = doc.add_paragraph()
    _para(p, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(CONTACT)
    _font(r, 8.55, MUTED)

    _heading(doc, "Profile")
    _text(doc, data["profile"], size=8.95)

    _heading(doc, "Role Match")
    for m in data.get("role_match", []):
        p = doc.add_paragraph()
        _para(p, after=1)
        r = p.add_run(f"{m['label']}: ")
        _font(r, 8.85, NAVY, bold=True)
        r = p.add_run(m["text"])
        _font(r, 8.85)

    _heading(doc, "Experience")
    for role in data.get("experience", []):
        p = doc.add_paragraph()
        _para(p, before=3, after=0)
        r = p.add_run(role["title"])
        _font(r, 9.75, NAVY, bold=True)
        r = p.add_run(f" - {role['company']}")
        _font(r, 9.75, bold=True)
        r = p.add_run(f" | {role['dates']}")
        _font(r, 9.2, MUTED)
        for b in role.get("bullets", []):
            _bullet(doc, b)

    _heading(doc, "Education")
    _text(doc, EDUCATION, size=9.1, bold=True, color=NAVY, after=1)
    if data.get("education_note"):
        _text(doc, data["education_note"], size=8.75)

    _heading(doc, "Tools & Keywords")
    _text(doc, data["keywords"], size=8.25, after=0)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = False
    doc.core_properties.author = "Matthew McCusker"
    doc.save(out_docx)
    return out_docx
